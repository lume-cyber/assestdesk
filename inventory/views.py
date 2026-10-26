import csv
import io
from datetime import date, datetime
from datetime import timedelta
from django.contrib import messages
from django.contrib.auth.models import User
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.core.paginator import Paginator
from django.db import transaction, IntegrityError
from django.db.models import Q
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.views.decorators.cache import never_cache
from .access import can_edit_finance, can_manage_lists, can_manage_users, permitted_assets, permitted_lists
from .forms import AssetForm, AssetListForm, FinanceForm, ManagedUserCreateForm, ManagedUserEditForm, MembershipForm
from .models import Asset, AssetEvent, AssetList, Membership

IMPORT_FIELDS = {
    'tag': ('tag', 'asset tag', 'asset_tag'), 'name': ('name', 'asset', 'asset name'),
    'category': ('category', 'type'), 'status': ('status',), 'manufacturer': ('manufacturer', 'brand'),
    'model': ('model',), 'serial_number': ('serial number', 'serial', 'serial_number'),
    'assigned_to': ('assigned to', 'assigned_to', 'owner'), 'location': ('location',),
    'purchase_date': ('purchase date', 'purchase_date'), 'warranty_expires': ('warranty expires', 'warranty_expires', 'warranty end'),
    'purchase_value': ('purchase value', 'purchase_value', 'value', 'cost'),
    'purchase_currency': ('purchase currency', 'purchase_currency', 'currency'),
    'notes': ('notes', 'description', 'comments'),
}

def _header(value):
    return ' '.join(str(value or '').strip().lower().replace('_',' ').split())

def _date_value(value):
    if value in (None, ''): return None
    if isinstance(value, datetime): return value.date()
    if isinstance(value, date): return value
    text = str(value).strip()
    for fmt in ('%Y-%m-%d','%d/%m/%Y','%m/%d/%Y','%d-%m-%Y'):
        try: return datetime.strptime(text, fmt).date()
        except ValueError: pass
    raise ValueError('use YYYY-MM-DD or a standard spreadsheet date')

def _decimal_value(value):
    if value in (None, ''): return None
    from decimal import Decimal, InvalidOperation
    try: result = Decimal(str(value).replace(',','').strip())
    except InvalidOperation: raise ValueError('purchase value must be a number')
    if result < 0: raise ValueError('purchase value cannot be negative')
    return result

def _rows_from_upload(upload):
    name = upload.name.lower()
    if name.endswith('.xlsx'):
        from openpyxl import load_workbook
        sheet = load_workbook(upload, read_only=True, data_only=True).active
        rows = sheet.iter_rows(values_only=True)
    elif name.endswith('.csv'):
        rows = csv.reader(io.TextIOWrapper(upload.file, encoding='utf-8-sig', newline=''))
    else:
        raise ValueError('upload an .xlsx or .csv file')
    rows = iter(rows)
    try: headers = next(rows)
    except StopIteration: raise ValueError('the file is empty')
    normalized = [_header(v) for v in headers]
    mapping = {}
    for field, aliases in IMPORT_FIELDS.items():
        for alias in aliases:
            if alias in normalized: mapping[field] = normalized.index(alias); break
    if 'tag' not in mapping or 'name' not in mapping:
        raise ValueError('the first row must include Tag and Asset/Name columns')
    return mapping, rows

def filtered(request):
    items = permitted_assets(request.user)
    q = request.GET.get('q','').strip()[:150]
    if q: items = items.filter(Q(name__icontains=q)|Q(tag__icontains=q)|Q(serial_number__icontains=q)|Q(assigned_to__icontains=q))
    list_id = request.GET.get('list','')
    if list_id:
        items = items.filter(asset_list_id=list_id) if list_id.isdigit() else items.none()
    status = request.GET.get('status','')
    if status: items = items.filter(status=status)
    return items

@login_required
@never_cache
def home(request):
    all_items = permitted_assets(request.user)
    today = timezone.localdate()
    params = request.GET.copy(); params.pop('page',None)
    return render(request,'inventory/home.html',{'page':Paginator(filtered(request),20).get_page(request.GET.get('page')),'lists':permitted_lists(request.user),'can_create':permitted_lists(request.user,True).exists(),'statuses':Asset.STATUSES,'query':params.urlencode(),'total':all_items.count(),'assigned':all_items.filter(status='assigned').count(),'available':all_items.filter(status='available').count(),'warranty':all_items.filter(warranty_expires__range=[today,today+timedelta(days=30)]).exclude(status='retired').count()})

@login_required
@never_cache
def detail(request,pk):
    asset = get_object_or_404(permitted_assets(request.user),pk=pk)
    return render(request,'inventory/detail.html',{'asset':asset,'can_edit':permitted_lists(request.user,True).filter(pk=asset.asset_list_id).exists(),'can_edit_finance':can_edit_finance(request.user),'events':asset.events.select_related('actor')[:30]})

@login_required
@never_cache
def asset_edit(request,pk=None):
    asset = get_object_or_404(permitted_assets(request.user,True),pk=pk) if pk else None
    if not permitted_lists(request.user,True).exists(): raise PermissionDenied
    previous_employee = asset.assigned_to if asset else ''
    form = AssetForm(request.POST or None,instance=asset,user=request.user)
    if request.method == 'POST' and form.is_valid():
        try:
            with transaction.atomic():
                saved = form.save()
                if saved.assigned_to != previous_employee:
                    saved.assignments.filter(returned_date__isnull=True).update(returned_date=timezone.localdate())
                    if saved.assigned_to:
                        AssignmentRecord.objects.create(asset=saved, employee=saved.assigned_to, assigned_date=timezone.localdate())
                changed = ', '.join(str(form.fields[f].label or f.replace('_',' ')) for f in form.changed_data)
                AssetEvent.objects.create(asset=saved,actor=request.user,action=('Updated: '+changed)[:250] if pk else 'Asset registered')
            messages.success(request,'Asset saved.')
            return redirect('detail',pk=saved.pk)
        except IntegrityError:
            form.add_error(None,'The asset could not be saved. Check that its tag is unique and try again.')
    return render(request,'inventory/form.html',{'form':form,'asset':asset})

@login_required
@never_cache
def import_assets(request):
    writable = permitted_lists(request.user, True)
    if not writable.exists(): raise PermissionDenied
    if request.method == 'POST':
        asset_list = get_object_or_404(writable, pk=request.POST.get('asset_list'))
        try:
            mapping, rows = _rows_from_upload(request.FILES.get('file'))
            created = 0
            with transaction.atomic():
                for row_number, row in enumerate(rows, start=2):
                    if not any(str(v or '').strip() for v in row): continue
                    def value(field):
                        index = mapping.get(field)
                        return row[index] if index is not None and index < len(row) else ''
                    tag, name = str(value('tag')).strip(), str(value('name')).strip()
                    if not tag or not name: raise ValueError(f'row {row_number}: Tag and Name are required')
                    category = str(value('category') or 'General products').strip()[:150]
                    status = str(value('status') or 'available').strip().lower().replace(' ', '_')
                    if status not in dict(Asset.STATUSES): status = 'available'
                    purchase_value = _decimal_value(value('purchase_value'))
                    purchase_currency = str(value('purchase_currency') or '').strip().upper()[:3]
                    if purchase_value is not None and not purchase_currency: raise ValueError(f'row {row_number}: Currency is required when Purchase value is supplied')
                    asset = Asset.objects.create(asset_list=asset_list, tag=tag, name=name, category=category, status=status, manufacturer=str(value('manufacturer') or '').strip(), model=str(value('model') or '').strip(), serial_number=str(value('serial_number') or '').strip(), assigned_to=str(value('assigned_to') or '').strip(), location=str(value('location') or '').strip(), purchase_date=_date_value(value('purchase_date')), purchase_value=purchase_value, purchase_currency=purchase_currency, warranty_expires=_date_value(value('warranty_expires')), notes=str(value('notes') or '').strip())
                    AssetEvent.objects.create(asset=asset, actor=request.user, action='Imported from spreadsheet')
                    created += 1
            messages.success(request, f'{created} asset(s) imported into {asset_list.name}.')
            return redirect('home')
        except Exception as exc:
            messages.error(request, f'Import stopped: {exc}')
    return render(request, 'inventory/import.html', {'lists': writable})

def csv_safe(value):
    text = '' if value is None else str(value)
    return "'"+text if text.lstrip().startswith(('=','+','-','@','\t','\r','\n')) else text

@login_required
@never_cache
def export_assets(request):
    response = HttpResponse(content_type='text/csv; charset=utf-8')
    response['Content-Disposition'] = 'attachment; filename="assets.csv"'
    writer = csv.writer(response)
    writer.writerow(['Tag','Asset','List','Category','Status','Serial number','Assigned to','Location','Purchase date','Purchase value','Purchase currency','Warranty expires'])
    for a in filtered(request).iterator():
        writer.writerow([csv_safe(v) for v in [a.tag,a.name,a.asset_list.name,a.category,a.get_status_display(),a.serial_number,a.assigned_to,a.location,a.purchase_date,a.purchase_value,a.purchase_currency,a.warranty_expires]])
    return response

@login_required
@never_cache
def asset_finance(request,pk):
    if not can_edit_finance(request.user): raise PermissionDenied
    asset=get_object_or_404(permitted_assets(request.user),pk=pk)
    form=FinanceForm(request.POST or None,instance=asset)
    if request.method=='POST' and form.is_valid():
        saved=form.save(); AssetEvent.objects.create(asset=saved,actor=request.user,action='Purchase value updated')
        messages.success(request,'Purchase information saved.'); return redirect('detail',pk=asset.pk)
    return render(request,'inventory/finance.html',{'form':form,'asset':asset})

@login_required
@never_cache
def lists(request):
    if not can_manage_lists(request.user): raise PermissionDenied
    form=AssetListForm(request.POST or None)
    if request.method=='POST' and form.is_valid():
        item=form.save(); messages.success(request,f'List “{item.name}” created.'); return redirect('list_manage',pk=item.pk)
    return render(request,'inventory/lists.html',{'form':form,'asset_lists':AssetList.objects.prefetch_related('memberships').all()})

@login_required
@never_cache
def list_manage(request,pk):
    if not can_manage_lists(request.user): raise PermissionDenied
    item=get_object_or_404(AssetList,pk=pk)
    details=AssetListForm(instance=item,prefix='details')
    member=MembershipForm(prefix='member')
    if request.method=='POST':
        action=request.POST.get('action')
        if action=='details':
            details=AssetListForm(request.POST,instance=item,prefix='details')
            if details.is_valid(): details.save(); messages.success(request,'List details saved.'); return redirect('list_manage',pk=pk)
        elif action=='member':
            member=MembershipForm(request.POST,prefix='member')
            if member.is_valid():
                data=member.cleaned_data; Membership.objects.update_or_create(asset_list=item,user=data['user'],defaults={'role':data['role']})
                messages.success(request,'List access saved.'); return redirect('list_manage',pk=pk)
        elif action=='remove':
            membership=get_object_or_404(item.memberships,pk=request.POST.get('membership'))
            membership.delete(); messages.success(request,'List access removed.'); return redirect('list_manage',pk=pk)
    return render(request,'inventory/list_manage.html',{'asset_list':item,'details_form':details,'member_form':member,'memberships':item.memberships.select_related('user').order_by('user__username')})

@login_required
@never_cache
def people(request):
    if not can_manage_users(request.user): raise PermissionDenied
    users=User.objects.select_related('profile').order_by('username')
    if not request.user.is_superuser: users=users.filter(is_superuser=False)
    return render(request,'inventory/people.html',{'managed_users':users})

@login_required
@never_cache
def user_create(request):
    if not can_manage_users(request.user): raise PermissionDenied
    form=ManagedUserCreateForm(request.POST or None,manager=request.user)
    if request.method=='POST' and form.is_valid():
        user=form.save(); messages.success(request,f'Account “{user.username}” created.'); return redirect('user_edit',pk=user.pk)
    return render(request,'inventory/user_form.html',{'form':form,'target':None})

@login_required
@never_cache
def user_edit(request,pk):
    if not can_manage_users(request.user): raise PermissionDenied
    target=get_object_or_404(User.objects.select_related('profile'),pk=pk)
    if target.is_superuser and not request.user.is_superuser: raise PermissionDenied
    form=ManagedUserEditForm(request.POST or None,manager=request.user,target=target)
    if request.method=='POST' and form.is_valid():
        form.save(); messages.success(request,'Account saved.'); return redirect('people')
    return render(request,'inventory/user_form.html',{'form':form,'target':target})

from .forms import PhotoForm, ServiceForm, AssignmentForm
from .models import AssignmentRecord
from django.http import Http404

@login_required
@never_cache
def asset_photo(request, pk):
    asset = get_object_or_404(permitted_assets(request.user), pk=pk)
    if not asset.photo_data: raise Http404
    response = HttpResponse(bytes(asset.photo_data), content_type='image/jpeg')
    response['Content-Security-Policy'] = "default-src 'none'"
    return response

@login_required
@never_cache
def asset_record(request, pk, kind, record_pk=None):
    asset = get_object_or_404(permitted_assets(request.user, True), pk=pk)
    classes = {'photo': PhotoForm, 'service': ServiceForm, 'assignment': AssignmentForm}
    if kind not in classes: raise Http404
    instance = None
    if record_pk:
        if kind != 'assignment': raise Http404
        instance = get_object_or_404(asset.assignments, pk=record_pk)
    kwargs = {'instance': instance} if kind != 'photo' else {}
    form = classes[kind](request.POST or None, request.FILES or None, **kwargs)
    if request.method == 'POST' and form.is_valid():
        with transaction.atomic():
            asset = get_object_or_404(permitted_assets(request.user, True).select_for_update(), pk=pk)
            if kind == 'photo':
                data = b''
                if form.cleaned_data['photo']:
                    from PIL import Image, ImageOps
                    photo = form.cleaned_data['photo']
                    photo.seek(0)
                    with Image.open(photo) as original:
                        picture = ImageOps.exif_transpose(original).convert('RGB')
                        picture.thumbnail((1600,1600))
                        output = io.BytesIO()
                        picture.save(output, format='JPEG', quality=85)
                        data = output.getvalue()
                asset.photo_data = data
                asset.save(update_fields=['photo_data'])
            else:
                record = form.save(commit=False)
                record.asset = asset
                if kind == 'assignment' and not record.returned_date and asset.assignments.filter(returned_date__isnull=True).exclude(pk=record.pk).exists():
                    form.add_error(None, 'Return the current assignment before adding another active assignment.')
                    return render(request, 'inventory/record.html', {'form':form,'asset':asset,'kind':kind})
                record.save()
                if kind == 'assignment':
                    current = asset.assignments.filter(returned_date__isnull=True).first()
                    asset.assigned_to = current.employee if current else ''
                    if current: asset.status = 'assigned'
                    elif asset.status == 'assigned': asset.status = 'available'
                    asset.save(update_fields=['assigned_to','status','updated_at'])
            AssetEvent.objects.create(asset=asset, actor=request.user, action=kind.title() + ' record updated')
        messages.success(request, 'Asset updated.')
        return redirect('detail', pk=pk)
    return render(request, 'inventory/record.html', {'form':form,'asset':asset,'kind':kind})
